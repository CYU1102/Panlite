"""Read-only independent extraction for explicitly listed local acceptance samples.

The output is a comparison reference, never certified ground truth. No Office
application, macros, document links, or embedded instructions are executed.
"""
import csv
import datetime
import io
import json
import sys
import zipfile
from pathlib import Path


def cell_text(value):
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, datetime.datetime):
        return value.date().isoformat() if value.time() == datetime.time() else value.isoformat()
    return str(value)


def table_cell(value):
    value = cell_text(value)
    return json.dumps(value, ensure_ascii=False) if any(char in value for char in '\t\n\r"\\') else value


def main():
    file = Path(sys.argv[1]).resolve()
    extension = file.suffix.lower()
    if file.stat().st_size > 100 * 1024 * 1024:
        raise ValueError("Reference extraction limited to 100 MB")
    if extension in (".docx", ".xlsx", ".pptx"):
        with zipfile.ZipFile(file) as archive:
            entries = archive.infolist()
            if len(entries) > 5000 or sum(entry.file_size for entry in entries) > 100 * 1024 * 1024:
                raise ValueError("Reference extraction ZIP budget exceeded")
    result = {"source": "independent-library", "groundTruth": False, "sections": [], "notes": []}
    if extension == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(str(file))
        result["source"] = "pypdf native text layer"
        result["pageCount"] = len(reader.pages)
        if len(reader.pages) > 100:
            raise ValueError("Reference extraction limited to 100 pages")
        result["sections"] = [{"pageNumber": index + 1, "content": page.extract_text() or ""} for index, page in enumerate(reader.pages)]
        result["notes"] = ["PDF text extraction may disagree on layout. Empty reference pages require image review; OCR is not ground truth."]
    elif extension == ".docx":
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph
        document = Document(str(file))
        lines = []
        for element in document.element.body:
            if element.tag.endswith("}p"):
                lines.append(Paragraph(element, document).text)
            elif element.tag.endswith("}tbl"):
                for row in Table(element, document).rows:
                    lines.append("\t".join(table_cell(cell.text) for cell in row.cells))
        result["source"] = "python-docx body paragraphs and tables"
        result["sections"] = [{"section": "正文", "content": "\n".join(lines)}]
        result["notes"] = ["Merged cells, revisions, text boxes, headers/footers and embedded drawings need manual review; repeated merged-cell values can differ."]
    elif extension == ".xlsx":
        from openpyxl import load_workbook
        workbook = load_workbook(file, read_only=True, data_only=True, keep_links=False)
        result["source"] = "openpyxl cached worksheet values"
        for worksheet in workbook:
            if worksheet.max_row * worksheet.max_column > 100000:
                raise ValueError("Reference worksheet limited to 100000 cells")
            lines = []
            source_rows = []
            for index, row in enumerate(worksheet.iter_rows(), 1):
                values = [cell.value for cell in row]
                while values and values[-1] is None:
                    values.pop()
                if any(value is not None for value in values):
                    lines.append("\t".join(table_cell(value) for value in values))
                    source_rows.append(index)
            result["sections"].append({"section": f"工作表：{worksheet.title}", "content": "\n".join(lines), "sourceRows": source_rows})
        workbook.close()
        result["notes"] = ["Cached values only; formulas are never recalculated. Percent formatting, uncached formulas, empty cells and merged-cell layouts can legitimately differ; inspect cell coordinates."]
    elif extension == ".pptx":
        from pptx import Presentation
        presentation = Presentation(str(file))
        result["source"] = "python-pptx slide text frames and tables"
        for index, slide in enumerate(presentation.slides, 1):
            lines = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    lines.extend(paragraph.text for paragraph in shape.text_frame.paragraphs)
                if shape.has_table:
                    for row in shape.table.rows:
                        lines.extend(cell.text for cell in row.cells)
            result["sections"].append({"pageNumber": index, "content": "\n".join(lines)})
        result["notes"] = ["Only visible slide text frames/tables; notes, group shapes, diagrams and reading order require manual review."]
    elif extension in (".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp"):
        from PIL import Image
        with Image.open(file) as image:
            result["image"] = {"width": image.width, "height": image.height, "mode": image.mode}
        result["source"] = "image metadata only"
        result["notes"] = ["No independent transcription: manual visual ground truth is required before measuring OCR accuracy."]
    elif extension in (".txt", ".md", ".markdown", ".csv", ".json", ".xml", ".log", ".yaml", ".yml", ".html", ".htm"):
        raw = file.read_bytes()
        if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
            encoding = "utf-16"
        else:
            try:
                raw.decode("utf-8-sig")
                encoding = "utf-8-sig"
            except UnicodeDecodeError:
                encoding = "gb18030"
        content = raw.decode(encoding)
        result["source"] = f"Python strict {encoding} decode"
        result["sections"] = [{"section": "正文", "content": content}]
        if extension == ".csv":
            rows = list(csv.reader(io.StringIO(content)))
            result["csv"] = {"rows": len(rows), "columnCounts": [len(row) for row in rows[:1000]]}
        result["notes"] = ["Encoding was inferred; non-BOM legacy encodings still require a human-readable check."]
    else:
        result["source"] = "unsupported reference format"
        result["notes"] = ["This format needs manual or converted-copy review."]
    if sum(len(section["content"]) for section in result["sections"]) > 4 * 1024 * 1024:
        raise ValueError("Reference text exceeds 4 MB budget")
    print(json.dumps(result, ensure_ascii=True))


if __name__ == "__main__":
    main()
